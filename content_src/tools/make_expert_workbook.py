"""One XLSX for the nurse-teacher: sheet1 priority questions needing an expert, sheet2 the rest (next update), sheet3 ticket answers."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
BANK = {'141': 'Hamshiralik ishi', '142': 'Patronaj', '138': 'Oliy hamshiralik'}
fs = json.load(open('content/final_status.json'))
import os
AI = {}
if os.path.exists('content/ai_suggest_results.jsonl'):
    for _l in open('content/ai_suggest_results.jsonl'):
        if _l.strip(): _r = json.loads(_l); AI[_r['id']] = _r
need = [r for r in fs if r['status'] in ('disputed', 'no_source')]
tickets = json.load(open('content/tickets_final.json'))
wb = Workbook()
def sheet(ws, header, rows, widths, dvcol=None):
    ws.append(header)
    for r in rows: ws.append(r)
    for c, w in zip('ABCDEFGHIJ', widths): ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=1):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F4C5C')
    ws.freeze_panes = 'C2'
    if dvcol and rows:
        dv = DataValidation(type='list', formula1='"to\'g\'ri,noto\'g\'ri"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'{dvcol}2:{dvcol}{len(rows) + 1}')
def row(r):
    o = r['options']; k = r.get('key') if isinstance(r.get('key'), int) else r.get('key_candidate')
    if not isinstance(k, int): k = r.get('file_key') if isinstance(r.get('file_key'), int) else None
    cand = f"{chr(65 + k)}) {o[k]}" if isinstance(k, int) and k < len(o) else ''
    note = (r.get('notes') or '').strip()
    a = AI.get(r['id'])
    if not cand and a and isinstance(a.get('ai_key'), int) and a['ai_key'] < len(o):
        cand = f"{chr(65 + a['ai_key'])}) {o[a['ai_key']]}  [AI taxmini, ishonch: {a['confidence']}]"
        note = (note + ' | AI: ' + a['why'] + (' (' + ','.join(a['flags']) + ')' if a.get('flags') else '')).strip(' |')
    return [r['id'], BANK[r['bank']], r['q'], '\n'.join(f"{chr(65 + i)}) {x}" for i, x in enumerate(o)), cand, note[:380], r['status'], '', '', '']
H = ['ID', 'Bank', 'Savol', 'Variantlar', 'Nomzod kalit', 'Qisqa izoh', 'Holat', "Mutaxassis: to'g'ri / noto'g'ri", "Mutaxassis: to'g'ri javob (A/B/C/D yoki matn)", 'Mutaxassis F.I.Sh. / sana']
W = (13, 16, 55, 60, 34, 55, 11, 20, 28, 24)
ws1 = wb.active; ws1.title = '1-versiya (ustuvor)'
sheet(ws1, H, [row(r) for r in need if r['priority'] or r['status'] == 'disputed'], W, 'H')
ws2 = wb.create_sheet('Keyingi yangilanish')
sheet(ws2, H, [row(r) for r in need if not (r['priority'] or r['status'] == 'disputed')], W, 'H')
ws3 = wb.create_sheet('Biletlar')
rows = []
for k, t in tickets.items():
    for it in t['items']:
        rows.append([it['id'], f"{({'1': 'Birinchi', '2': 'Ikkinchi', 'oliy': 'Oliy'})[t['cat']]} toifa, bilet № {t['number']}", it['question'], '\n'.join(f'{i + 1}. {p}' for i, p in enumerate(it['points'])), it['source'], ', '.join(it['flags']) + ((' | ' + it['note']) if it['note'] else ''), '', '', '', ''])
sheet(ws3, ['ID', 'Bilet', 'Savol', 'Namuna-javob (3–6 band)', 'Manba', 'Belgilar (tekshirish kerak)', '', "Mutaxassis: to'g'ri / noto'g'ri", "Mutaxassis: tuzatilgan javob", 'Mutaxassis F.I.Sh. / sana'], rows, (16, 24, 50, 70, 50, 38, 3, 20, 50, 24), 'H')
out = sys.argv[1] if len(sys.argv) > 1 else 'content/mutaxassis_jadvali.xlsx'
wb.save(out); print('sheet1', ws1.max_row - 1, 'sheet2', ws2.max_row - 1, 'sheet3', ws3.max_row - 1, '->', out)
