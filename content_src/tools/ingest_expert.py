"""Read the filled specialist XLSX -> content/expert_results.json (status 'expert')."""
import json, sys, re
from openpyxl import load_workbook
ws = load_workbook(sys.argv[1]).active
st = {r['id']: r for r in json.load(open('content/question_status.json'))}
out = []; skipped = []
for row in list(ws.iter_rows(values_only=True))[1:]:
    i, verdict, ans, who = row[0], (row[5] or '').strip().lower(), (row[6] or '').strip(), (row[7] or '').strip()
    if not i or not verdict: continue
    r = st[i]; opts = r['options']; key = None
    if verdict.startswith("to'g'ri"):
        key = r.get('key') if isinstance(r.get('key'), int) else r.get('key_candidate')
    elif ans:
        m = re.match(r'^\s*([A-Ea-e])\b', ans)
        if m: key = ord(m.group(1).upper()) - 65
        else:
            for k, o in enumerate(opts):
                if o.strip().lower() == ans.lower(): key = k
    if not isinstance(key, int) or key >= len(opts) or not who:
        skipped.append(i); continue
    out.append({'id': i, 'status': 'expert', 'key': key, 'source': f'Mutaxassis tekshiruvi, {who}', 'source_ref': {'file_or_url': 'mutaxassis_jadvali.xlsx', 'page_or_clause': i}})
json.dump(out, open('content/expert_results.json', 'w'), ensure_ascii=False, indent=0)
print(len(out), 'expert,', len(skipped), 'skipped (no key or no F.I.Sh.):', skipped[:10])
